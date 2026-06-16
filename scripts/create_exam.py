import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Add a Title
title = doc.add_heading('Avaliação de Estrutura de Dados', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Nome do Aluno: ________________________________________________________\nData: ___/___/______')

# Intro
doc.add_heading('Instruções:', level=2)
doc.add_paragraph('1. Esta prova contém 4 questões (2 teóricas e 2 práticas).')
doc.add_paragraph('2. Responda as questões de código utilizando preferencialmente Python, pseudocódigo ou a linguagem estudada em sala de aula, garantindo a correta lógica de manipulação das estruturas.')
doc.add_paragraph('3. Justifique suas respostas teóricas detalhadamente.')

# Question 1
doc.add_heading('Questão 1 (Teórica) - Análise de Cenários e Estruturas de Dados', level=2)
p = doc.add_paragraph()
p.add_run('Contexto: ').bold = True
p.add_run('Imagine que você está projetando a arquitetura principal de um editor de código (como o VS Code). O editor precisa gerenciar várias frentes de dados simultaneamente:\n')
p.add_run('a) O histórico de edições do usuário para permitir que ele desfaça (Ctrl+Z) as últimas ações.\n')
p.add_run('b) Uma fila de requisições de linting (análise de código em segundo plano) onde os arquivos recém modificados entram para serem analisados na ordem em que foram alterados.\n')
p.add_run('c) O armazenamento em memória do texto de um arquivo aberto, cujo tamanho cresce conforme o usuário digita novas linhas, exigindo acessos rápidos pelo número da linha.\n')

p = doc.add_paragraph()
p.add_run('Pergunta: ').bold = True
p.add_run('Com base nas estruturas de dados (Pilhas, Filas, Listas Dinâmicas/Dynamic Arrays), indique qual é a estrutura ideal para gerenciar CADA UM dos itens acima (a, b e c). Justifique tecnicamente sua escolha com base na forma de inserção e acesso dos dados e nas propriedades de cada estrutura estudada.')

# Question 2
doc.add_heading('Questão 2 (Teórica) - Listas Encadeadas vs Listas Dinâmicas', level=2)
p = doc.add_paragraph()
p.add_run('Contexto: ').bold = True
p.add_run('Em um aplicativo de streaming de áudio, o usuário criou uma "Playlist Infinita". A regra dessa playlist é: ao terminar a última música, ela recomeça automaticamente da primeira. O usuário deseja avançar rapidamente para a próxima música e voltar para a anterior sem atrasos. Durante o uso, o sistema precisa frequentemente inserir e remover músicas do meio desta playlist por conta de propagandas e recomendações dinâmicas.\n')

p = doc.add_paragraph()
p.add_run('Pergunta: ').bold = True
p.add_run('Qual variação específica de Lista Encadeada (Simples, Dupla ou Circular) é a ideal para suportar todas essas características? Explique detalhadamente como a mecânica de seus ponteiros satisfaz os requisitos (recomeçar ao final, avançar/retroceder). Além disso, explique por que o uso de uma Lista Dinâmica (Dynamic Array) causaria um problema de performance no cenário de inserções e remoções no meio da playlist.')


# Question 3
doc.add_heading('Questão 3 (Prática) - Inversão Parcial com Pilhas e Filas', level=2)
p = doc.add_paragraph()
p.add_run('Contexto: ').bold = True
p.add_run('Muitas vezes precisamos reorganizar buffers de dados em que apenas uma parte da mensagem chegou fora de ordem.\n')

p = doc.add_paragraph()
p.add_run('Pergunta: ').bold = True
p.add_run('Escreva uma função/método que receba uma Fila (Queue) e um número inteiro k. A função deve inverter a ordem apenas dos primeiros k elementos da fila, deixando a ordem dos demais elementos inalterada.\n')
p.add_run('Restrições:\n')
p.add_run('- Você pode usar apenas UMA Pilha (Stack) como estrutura de dados auxiliar.\n')
p.add_run('- Utilize apenas as operações básicas da Fila (enqueue, dequeue, is_empty, size/front) e da Pilha (push, pop, is_empty, top). Não acesse índices internos da estrutura.')


# Question 4
doc.add_heading('Questão 4 (Prática) - Manipulação de Ponteiros em Lista Duplamente Encadeada', level=2)
p = doc.add_paragraph()
p.add_run('Contexto: ').bold = True
p.add_run('Em um histórico de navegação implementado como uma Lista Duplamente Encadeada, cada nó possui os atributos `url`, `prev` (anterior) e `next` (próximo). Um bug no sistema fez com que cliques repetidos na mesma página gerassem nós consecutivos duplicados no histórico (Ex: Página A -> Página B -> Página B -> Página C -> Página B).\n')

p = doc.add_paragraph()
p.add_run('Pergunta: ').bold = True
p.add_run('Dado um ponteiro para a `head` (cabeça) dessa Lista Duplamente Encadeada, escreva uma função para percorrer a lista e remover apenas as duplicatas CONSECUTIVAS. No exemplo acima, a lista resultante deve ser A -> B -> C -> B. Modifique os ponteiros `next` e `prev` adequadamente e não se esqueça de tratar os casos onde a duplicata pode ser o último nó da lista. Você não pode criar uma lista nova, a operação deve ocorrer in-place.')


# Page Break for Answer Key
doc.add_page_break()

# Gabarito
doc.add_heading('GABARITO', 0)

# Gabarito Q1
doc.add_heading('Respostas - Questão 1:', level=3)
doc.add_paragraph('a) Pilha (Stack): A funcionalidade de "desfazer" opera no princípio LIFO (Last-In, First-Out). A última edição feita pelo usuário deve ser a primeira a ser desfeita. A inserção (push) e remoção (pop) no topo da pilha mapeiam perfeitamente esse comportamento de empilhar e desempilhar estados ou edições de código.')
doc.add_paragraph('b) Fila (Queue): O agendamento de tarefas em segundo plano deve seguir o princípio FIFO (First-In, First-Out), de modo que o primeiro arquivo modificado seja analisado primeiro para garantir justiça (fairness). Requisições entram no fim da fila (enqueue) e são consumidas no início (dequeue).')
doc.add_paragraph('c) Lista Dinâmica (Dynamic Array): Linhas de texto exigem acesso rápido e aleatório em tempo constante (O(1)) pelo índice da linha (ex: ir para a linha 150), algo que as listas encadeadas não suportam bem (são O(N)). Além disso, as listas dinâmicas ajustam seu tamanho conforme necessário na memória (via realocação), encaixando-se perfeitamente em armazenar e acessar linhas consecutivas de forma rápida.')

# Gabarito Q2
doc.add_heading('Respostas - Questão 2:', level=3)
doc.add_paragraph('A estrutura ideal é uma Lista Duplamente Encadeada Circular.\n'
'- Circular: Porque o ponteiro "next" do último nó aponta novamente para a cabeça da lista (e o "prev" do primeiro nó aponta para o final). Isso resolve automaticamente a exigência de "recomeçar a playlist ao terminar a última música", criando um loop contínuo.\n'
'- Duplamente Encadeada: Por ter um ponteiro "prev" além do "next", permite retroceder (voltar à música anterior) de forma natural com complexidade O(1) partindo do nó atual.\n\n'
'Problema da Lista Dinâmica (Dynamic Array): Inserir ou remover elementos no meio de uma lista dinâmica custa O(N) em tempo, pois exige o deslocamento (shift) de todos os elementos subsequentes na memória para abrir espaço (na inserção) ou fechar o buraco (na remoção). Na Lista Encadeada, caso já tenhamos o ponteiro para o local da playlist, inserções e remoções exigem apenas a reatribuição de ponteiros, com complexidade O(1), sendo muito mais eficiente para essa dinâmica.')

# Gabarito Q3
doc.add_heading('Respostas - Questão 3:', level=3)
code3 = """def inverter_k_elementos(fila, k):
    # Condições de guarda
    if fila.is_empty() or k <= 0:
        return
        
    pilha = Pilha()
    
    # 1. Desenfileira os primeiros k elementos e empilha
    for _ in range(k):
        pilha.push(fila.dequeue())
        
    # 2. Desempilha os elementos e coloca de volta na fila
    # Isso fará com que os primeiros k entrem na fila de forma invertida,
    # porém eles irão parar no final da fila.
    while not pilha.is_empty():
        fila.enqueue(pilha.pop())
        
    # 3. Restaura a ordem girando os elementos não-invertidos.
    # Os elementos originais restantes devem ser removidos do início 
    # e colocados no final.
    elementos_restantes = fila.size() - k
    for _ in range(elementos_restantes):
        fila.enqueue(fila.dequeue())
"""
doc.add_paragraph(code3, style='Intense Quote')
doc.add_paragraph('Lógica:\n1. Usamos as propriedades LIFO da pilha para inverter os K primeiros elementos.\n2. Como enfileiramos eles de volta no final da fila, a ordem geral ficou errada (os originais que não deviam mudar ficaram na frente).\n3. Rotacionamos o restante da fila ("tamanho - k" vezes), tirando do início e colocando no final, assim os elementos invertidos voltam para a frente.')


# Gabarito Q4
doc.add_heading('Respostas - Questão 4:', level=3)
code4 = """def remover_duplicatas_consecutivas(head):
    if head is None:
        return None
        
    atual = head
    
    while atual is not None and atual.next is not None:
        # Se encontrou duplicata consecutiva
        if atual.url == atual.next.url:
            duplicado = atual.next
            
            # Pula o nó duplicado atualizando o ponteiro next
            atual.next = duplicado.next
            
            # Se o próximo após o duplicado existir, o prev dele
            # deve apontar para o nó 'atual'
            if duplicado.next is not None:
                duplicado.next.prev = atual
                
            # Obs: Em linguagens sem garbage collection automático (ex: C), 
            # é necessário desalocar a memória de 'duplicado' (free).
            
            # Não avançamos o ponteiro 'atual' para o 'atual.next' ainda,
            # pois pode haver mais de uma duplicata consecutiva da mesma URL 
            # (Ex: A -> B -> B -> B). O laço testará o mesmo 'atual' 
            # contra o novo 'atual.next'.
        else:
            atual = atual.next
            
    return head
"""
doc.add_paragraph(code4, style='Intense Quote')
doc.add_paragraph('Lógica: O algoritmo navega com o ponteiro "atual". Sempre que "atual.url" é igual a "atual.next.url", os ponteiros do nó "atual" e do nó "atual.next.next" são remapeados, excluindo "atual.next" da malha. É vital a checagem `duplicado.next is not None` para evitar erros quando a duplicata é o último elemento da lista. Importante também notar que, em caso de remoção, o "atual" não avança para continuarmos tratando sequências longas sem ter que voltar atrás.')

doc.save('Prova_Estruturas_de_Dados.docx')
